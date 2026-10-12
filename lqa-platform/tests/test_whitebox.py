# -*- coding: utf-8 -*-
"""白盒测试：按分支/边界构造，覆盖未命中路径并暴露缺陷。

运行：python tests/test_whitebox.py
约定：
  - test_*        → 契约测试，应全部通过
  - test_probe_*  → 规格探针，断言「文档承诺的正确行为」；其失败即代表当前已确认的缺陷
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from docx import Document
from openpyxl import Workbook
from fastapi.testclient import TestClient
import pandas as pd

import api_server
from lqa import exporters as ex
from lqa import generate as gen_mod
from lqa import llm as llm_mod
from lqa import skills as sk
from lqa.analysis import split_requirements
from lqa.config import Settings, load_env_file, load_settings
from lqa.exporters import LOC_COLUMNS, cases_to_dataframe, save_output, to_xlsx_bytes
from lqa.llm import LLMError, build_system_prompt, chat_json, extract_json_payload, load_few_shots
from lqa.parsers import merge_sources, parse_docx, parse_md, parse_txt, parse_upload, parse_xlsx
from lqa.quality import dedupe, gate_report, merge_cases, normalize_case, score_case
from lqa.rules import build_rule_cases, _extract_samples, _pick_categories, _priority_for, _preconditions

client = TestClient(api_server.app)


# ---------------- 构造器 ----------------
def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("商城需求说明")
    doc.add_paragraph("")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "模块"
    t.cell(0, 1).text = "说明"
    t.cell(1, 0).text = "商城"
    t.cell(1, 1).text = "多行\n文本"
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def _xlsx_bytes(rows: int = 3, sheets=("S1", "S2")) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(name)
        for r in range(1, rows + 1):
            ws.cell(r, 1, "r%d" % r)
            ws.cell(r, 2, None if r == 2 else "v%d" % r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


GOOD = {"module": "商城", "feature": "限时礼包标题溢出检查", "category": "TEXT_OVERFLOW",
        "priority": "P0", "source_text": "Limited Offer", "target_text": "Zeitlich begrenztes Angebot",
        "preconditions": "语言 de-DE，账号可进入商城", "steps": "1. 切换语言\n2. 进入商城\n3. 观察标题是否截断",
        "expected": "标题完整显示，无截断/省略号，不压价格控件", "notes": "德语长文案高风险"}


# ---------------- parsers ----------------
def test_parsers_all_formats_and_errors():
    assert parse_txt("a  \r\nb\r\n\r\n\r\nc ") == "a\nb\n\nc"
    assert parse_md(b"# T\n\nbody") == "# T\n\nbody"
    assert "商城需求说明" in parse_docx(_docx_bytes())
    assert "[表格1]" in parse_docx(_docx_bytes())
    assert "模块 | 说明" in parse_docx(_docx_bytes())
    assert "多行 文本" in parse_docx(_docx_bytes())
    tx = parse_xlsx(_xlsx_bytes())
    assert "[Sheet] S1" in tx and "[Sheet] S2" in tx and "r1 | v1" in tx
    assert "r520" not in parse_xlsx(_xlsx_bytes(rows=520, sheets=("Big",)))
    assert "hi" in parse_upload("a.markdown", b"hi")
    assert "hi" in parse_upload("a.text", b"hi")
    assert "hi" in parse_upload("a.log", b"hi")
    assert "hi" in parse_upload("A.TXT", b"hi")
    assert "商城" in parse_upload("a.docx", _docx_bytes())
    assert "r1" in parse_upload("a.xlsm", _xlsx_bytes())
    for bad in ("a.doc", "a.xls", "a.pdf", "noext"):
        try:
            parse_upload(bad, b"x")
            raise AssertionError("应拒绝: " + bad)
        except ValueError:
            pass


def test_parsers_merge_branches():
    assert merge_sources("A", "B") == "A\n\nB"
    assert merge_sources("", "B") == "B"
    assert merge_sources("A", "") == "A"
    assert merge_sources("", "") == ""


# ---------------- config ----------------
def test_config_env_file_and_masking():
    tmp = Path(tempfile.mkdtemp())
    env = tmp / ".env"
    env.write_text('# comment\n\nLQA_WB_A="quoted value"\nLQA_WB_B = bare \n=novalue\n', encoding="utf-8")
    for k in ("LQA_WB_A", "LQA_WB_B"):
        os.environ.pop(k, None)
    load_env_file(env)
    assert os.environ.get("LQA_WB_A") == "quoted value"
    assert os.environ.get("LQA_WB_B") == "bare"
    load_env_file(tmp / "missing.env")
    assert Settings("deepseek", "u", "", "m").masked_key() == ""
    assert Settings("deepseek", "u", "short", "m").masked_key() == "已配置"
    assert Settings("deepseek", "u", "sk-abcdefghijklmno", "m").masked_key() == "sk-a…lmno"
    d = Settings("custom", "http://x", "k", "m").public_dict()
    assert d["has_api_key"] is True and d["api_key_masked"] == "已配置" and "temperature" in d


def test_config_provider_branches():
    for k in ("LQA_MODEL", "LQA_BASE_URL", "LQA_API_KEY"):
        os.environ.pop(k, None)
    for p in ("deepseek", "openai", "mimo", "custom"):
        s = load_settings(provider=p, api_key="k", model="m")
        assert s.provider == p and s.base_url.startswith("http")
    assert load_settings(provider="mimo", api_key="k").model == "mimo-v2.5"
    assert load_settings(provider="openai", api_key="k").model == "gpt-4o-mini"
    assert load_settings(provider="custom", api_key="k").model == "qwen2.5"
    b = load_settings(provider="bogus", api_key="k")
    assert b.provider == "deepseek" and b.model == "deepseek-chat"
    assert load_settings(provider="deepseek", base_url="http://x/y/", api_key="k").base_url == "http://x/y"


# ---------------- exporters ----------------
def test_export_helpers_branches():
    assert ex._numbered("") == ""
    assert ex._numbered("1. a\n2. b") == "1.a\n2.b"
    assert ex._numbered("a；b") == "1.a\n2.b"
    assert ex._numbered("a\n\n b ") == "1.a\n2.b"
    assert ex._split_module("") == ("", "")
    assert ex._split_module("Main") == ("Main", "")
    assert ex._split_module("Main/Sub") == ("Main", "Sub")
    assert ex._split_module("Main\\Sub/Deep") == ("Main", "Sub/Deep")
    assert ex._localization_checks("CULTURAL").count("\n") >= 4
    assert ex._localization_checks("未知类别").count("\n") == 3
    assert ex._test_statement({"feature": "F"}) == "F"
    assert "校验" in ex._test_statement({"module": "A/B", "category": "OTHER"})


def test_export_dataframe_and_save():
    df = cases_to_dataframe([{"module": "A/B", "feature": "", "category": "OTHER"}])
    assert list(df.columns) == LOC_COLUMNS
    row = df.iloc[0]
    assert row["用例编号"] == "LOC-001" and row["主模块"] == "A" and row["子模块"] == "B"
    assert row["入口截图 / 测试截图"] == "" and row["实际结果"] == "" and row["测试人"] == ""
    assert row["预期结果"].startswith("1.")
    assert cases_to_dataframe([], id_prefix="abc").empty
    assert to_xlsx_bytes([dict(GOOD)], sheet_name="x" * 40)[:2] == b"PK"
    tmp = Path(tempfile.mkdtemp())
    folder = save_output([dict(GOOD)], name='非法<>:"/\\|?*名称', out_root=tmp)
    assert folder.exists() and (folder / "generation_meta.json").exists()
    assert (folder / "本地化测试用例_非法名称.xlsx").exists()
    f2 = save_output([dict(GOOD)], name="R", out_root=tmp,
                     reports={"coverage_markdown": "# C", "gap_review_summary": "", "quality_markdown": "# Q"})
    assert (f2 / "覆盖率报告.md").read_text(encoding="utf-8") == "# C"
    assert not (f2 / "需求缺口审查.md").exists()
    assert (f2 / "质量门禁.md").exists()
    ex._write_if("# x", f2 / "generation_meta.json" / "x.md")  # OSError 分支：不抛即通过
    ex._write_if(None, f2 / "never.md")


# ---------------- quality ----------------
def test_quality_normalize_branches():
    assert normalize_case("not a dict", source_lang="en-US", target_lang="de-DE") is None
    n = normalize_case(dict(GOOD), source_lang="en-US", target_lang="de-DE")
    assert n["category"] == "TEXT_OVERFLOW" and n["priority"] == "P0"
    assert n["steps"].count("\n") == 2
    assert n["status"] == "待审" and n["actual"] == ""
    for miss in ("module", "feature", "steps", "expected"):
        bad = dict(GOOD); bad[miss] = ""
        assert normalize_case(bad, source_lang="en-US", target_lang="de-DE") is None, miss
    v = dict(GOOD); v["expected"] = "正常显示"
    assert normalize_case(v, source_lang="en-US", target_lang="de-DE") is None
    p = dict(GOOD); p["priority"] = "X9"
    assert normalize_case(p, source_lang="en-US", target_lang="de-DE")["priority"] == "P1"
    c = dict(GOOD); c["category"] = "赛季结算"
    assert normalize_case(c, source_lang="en-US", target_lang="de-DE")["category"] == "赛季结算"
    zh = dict(GOOD); zh["category"] = "文本溢出/截断"
    assert normalize_case(zh, source_lang="en-US", target_lang="de-DE")["category"] == "TEXT_OVERFLOW"
    lo = {"module": "M", "feature": "F", "category": "OTHER", "priority": "P1", "steps": "a\nb", "expected": "x" * 30}
    n2 = normalize_case(lo, source_lang="en-US", target_lang="de-DE")
    assert n2["source_lang"] == "en-US" and n2["target_lang"] == "de-DE"


def test_quality_score_and_gate_branches():
    base = dict(GOOD, source="ai")
    s0 = score_case(base)
    assert s0 > 40
    pen = dict(base); pen["steps"] = "1. 检查要点：\n2. 截图记录实际表现"
    assert score_case(pen) < s0
    vague = dict(base); vague["expected"] = "正常"
    assert score_case(vague) < s0
    rule = {"module": "M", "feature": "F", "category": "CULTURAL", "priority": "P0", "source": "rules",
            "steps": "1. a\n2. b\n3. c", "expected": "x" * 40, "source_text": "", "target_text": "",
            "preconditions": "", "notes": "规则打底"}
    assert score_case(rule) <= 72
    assert 0 <= score_case({}) <= 100
    assert score_case(dict(base, category="PLACEHOLDER", steps="1. {amount} 被替换", expected="{amount} 占位已替换，无字面量")) > 0
    rep = gate_report([])
    assert rep["count"] == 0 and rep["avg_score"] == 0 and rep["min_score"] == 0 and rep["pass"] is True
    bad = {"module": "M", "feature": "", "category": "OTHER", "priority": "P1", "steps": "no number",
           "expected": "", "quality_score": 10, "source": "rules"}
    r2 = gate_report([bad])
    assert r2["issue_count"] >= 1 and r2["pass"] is False
    custom = {"module": "M", "feature": "abcdef", "category": "赛季结算", "priority": "P1",
              "steps": "1. a", "expected": "完整显示，无截断", "quality_score": 80, "source": "rules"}
    assert all("非法" not in i for i in gate_report([custom])["issues"])
    one = [{"module": "M", "feature": "F", "category": "X", "source_text": "s", "expected": "e"}]
    assert len(dedupe(one + one)) == 1
    assert len(merge_cases(one, one, count=None)) == 1
    two = one + [dict(one[0], feature="G", expected="e2")]
    assert len(merge_cases(one, two[1:], count=None)) == 2


# ---------------- generate ----------------
def test_generate_balanced_top_and_paths():
    assert gen_mod._balanced_top([], 5) == []
    cs = [{"module": "M%d" % i, "category": "c", "quality_score": i} for i in range(6)]
    assert len(gen_mod._balanced_top(cs, None)) == 6
    assert len(gen_mod._balanced_top(cs, 3)) == 3
    a = split_requirements("【商城系统】\n1. 商店展示限时礼包标题与价格\n2. 售罄商品置灰显示已售罄\n")
    s = Settings("custom", "http://x", "sk-test-1234567890", "m")
    r1 = gen_mod.generate(s, a, source_lang="en-US", target_lang="de-DE", modules=a.modules,
                          categories=["TEXT_OVERFLOW"], extra="", use_ai=False, count=0)
    assert r1["generator"] == "rules" and r1["cases"]
    try:
        gen_mod.generate(Settings("custom", "http://x", "", "m"), a, source_lang="en-US", target_lang="de-DE",
                         modules=a.modules, categories=[], extra="", use_ai=True, count=2)
        raise AssertionError("无 Key + use_ai 应抛 LLMError")
    except LLMError:
        pass
    orig = gen_mod.chat_json
    try:
        gen_mod.chat_json = lambda *aa, **kk: {"cases": [dict(GOOD)]}
        r2 = gen_mod.generate(s, a, source_lang="en-US", target_lang="de-DE", modules=a.modules,
                              categories=["TEXT_OVERFLOW"], extra="", use_ai=True, count=4)
        assert "llm" in r2["generator"]
        def _boom(*aa, **kk):
            raise LLMError("boom")
        gen_mod.chat_json = _boom
        r3 = gen_mod.generate(s, a, source_lang="en-US", target_lang="de-DE", modules=a.modules,
                              categories=["TEXT_OVERFLOW"], extra="", use_ai=True, count=4)
        assert r3["generator"] == "rules" and "boom" in r3["ai_error"]
        gen_mod.chat_json = lambda *aa, **kk: {"cases": [{"module": "M"}]}
        r4 = gen_mod.generate(s, a, source_lang="en-US", target_lang="de-DE", modules=a.modules,
                              categories=["TEXT_OVERFLOW"], extra="", use_ai=True, count=4)
        assert r4["ai_error"]
    finally:
        gen_mod.chat_json = orig


def test_generate_user_prompt_branches():
    a = split_requirements("【商城】\n1. 商店展示\n")
    p = gen_mod.build_user_prompt(a, count=3, modules=a.modules, categories=["TEXT_OVERFLOW"],
                                  source_lang="en-US", target_lang="de-DE", extra="优先商店",
                                  skeleton=[{"module": "M", "feature": "F"}],
                                  selected_skills=[{"name": "赛季要点", "body": "结算边界检查", "description": "d"}])
    assert "赛季要点" in p and "结算边界检查" in p and "优先商店" in p
    p2 = gen_mod.build_user_prompt(a, count=3, modules=[], categories=[], source_lang="en-US",
                                   target_lang="de-DE", extra="", skeleton=[], selected_skills=None)
    assert "skeleton_cases" in p2


# ---------------- llm ----------------
def test_llm_extract_json_variants():
    assert extract_json_payload('{"cases":[1]}')["cases"] == [1]
    assert extract_json_payload('[{"a":1}]')["cases"][0]["a"] == 1
    assert extract_json_payload('```json\n{"cases":[]}\n```')["cases"] == []
    assert extract_json_payload('说明文字 {"cases":[2]} 结尾')["cases"] == [2]
    for bad in ("", "no json here"):
        try:
            extract_json_payload(bad)
            raise AssertionError("应抛 LLMError: " + repr(bad))
        except LLMError:
            pass


def test_llm_few_shots_and_prompt():
    tmp = Path(tempfile.mkdtemp())
    pd.DataFrame([{"模块": "商城", "功能点": "标题溢出", "步骤": "1. 切换语言\n2. 观察标题是否截断",
                   "预期结果": "完整显示，无截断", "类别": "TEXT_OVERFLOW", "优先级": "P0", "备注": "n"}
                  ]).to_excel(tmp / "old.xlsx", index=False)
    pd.DataFrame([{"主模块": "邮件", "子模块": "补偿", "测试说明": "占位符替换",
                   "复现路径": "1. 打开邮件\n2. 查看正文占位符是否被替换", "预期结果": "显示实际数量",
                   "前置条件": "语言 de-DE", "备注": "n"}]).to_excel(tmp / "new.xlsx", index=False)
    pd.DataFrame([{"foo": "bar"}]).to_excel(tmp / "bad.xlsx", index=False)
    shots = load_few_shots(tmp, limit=12)
    mods = [str(s.get("module") or "") for s in shots]
    assert any("商城" in m for m in mods)
    assert any("邮件" in m for m in mods)
    assert load_few_shots(tmp / "nope", limit=3)
    assert "LQA" in build_system_prompt(None)
    withsk = build_system_prompt(None, selected_skills=[{"name": "N", "description": "", "body": "B"}])
    assert "领域测试知识包" in withsk


class _FakeResp:
    def __init__(self, content):
        self.choices = [type("C", (), {"message": type("M", (), {"content": content})()})()]


class _FakeClient:
    def __init__(self, seq):
        self.seq = list(seq)
        self.calls = []
        comps = self

        def create(**kw):
            comps.calls.append(kw)
            item = comps.seq.pop(0)
            if isinstance(item, Exception):
                raise item
            return _FakeResp(item)
        self.chat = type("Chat", (), {"completions": type("Cp", (), {"create": staticmethod(create)})()})()


def test_llm_chat_json_paths():
    s = Settings("custom", "http://x", "sk-test-1234567890", "m")
    orig = llm_mod._client
    try:
        llm_mod._client = lambda st: _FakeClient(['{"cases":[]}'])
        assert chat_json(s, "p", system="sys") == {"cases": []}
        llm_mod._client = lambda st: _FakeClient([Exception("400 invalid response_format"), '{"cases":[1]}'])
        assert chat_json(s, "p", system="sys")["cases"] == [1]
        llm_mod._client = lambda st: _FakeClient([""])
        try:
            chat_json(s, "p", system="sys")
            raise AssertionError("空 content 应抛 LLMError")
        except LLMError:
            pass
        llm_mod._client = lambda st: _FakeClient([Exception("timeout")])
        try:
            chat_json(s, "p", system="sys")
            raise AssertionError("网络错误应抛 LLMError")
        except LLMError:
            pass
        llm_mod._client = orig
        try:
            chat_json(Settings("custom", "http://x", "", "m"), "p")
            raise AssertionError("无 Key 应抛 LLMError")
        except LLMError:
            pass
    finally:
        llm_mod._client = orig


# ---------------- skills ----------------
def test_skills_error_paths_and_helpers():
    tmp = Path(tempfile.mkdtemp())
    assert sk.list_skills(tmp) == []
    assert sk.delete_skill("nope", root=tmp) is False
    assert sk.resolve_skills(["nope", "..", "///"], root=tmp) == []
    assert sk.render_skills_block([]) == ""
    assert "N" in sk.render_skills_block([{"name": "N", "body": "B", "description": "D"}])
    for fn, exc in ((lambda: sk.get_skill("nope", root=tmp), KeyError),
                    (lambda: sk.update_skill("nope", name="x", root=tmp), KeyError)):
        try:
            fn()
            raise AssertionError("应抛 " + exc.__name__)
        except exc:
            pass
    for bad in ("..", "/", "///", "   ", "???"):
        try:
            sk._ensure_safe_id(bad)
            raise AssertionError("应拒绝 " + repr(bad))
        except ValueError:
            pass
    (tmp / "broken.json").write_text("{ 坏", encoding="utf-8")
    assert sk.list_skills(tmp) == []
    assert sk.skill_id("../../etc/passwd") == "etc-passwd"
    assert sk.skill_id("赛季通行证") == ""


# ---------------- rules ----------------
def test_rules_helpers_branches():
    assert _pick_categories("按钮文案溢出", [])
    assert _pick_categories("宗教敏感内容", ["CULTURAL"]) == ["CULTURAL"]
    assert _pick_categories("完全无关的一句话", ["TEXT_OVERFLOW"]) == ["TEXT_OVERFLOW"]
    # 优先级契约（见 _priority_for docstring）：按需求重要度分级，非按类别一刀切
    assert _priority_for("OTHER", "登录充值") == "P0"
    assert _priority_for("OTHER", "含宗教敏感元素") == "P0"
    assert _priority_for("TERMINOLOGY", "随便") == "P1"
    assert _priority_for("ENCODING", "随便") == "P1"
    assert _priority_for("PLACEHOLDER", "随便") == "P1"
    assert _priority_for("CULTURAL", "随便") == "P1"
    assert _priority_for("TEXT_OVERFLOW", "随便") == "P2"
    assert _priority_for("LAYOUT", "随便") == "P2"
    assert _priority_for("OTHER", "随便一句") == "P2"
    assert "Hello" in _extract_samples('文案 “Hello World” 需要替换')[0]
    assert "登录" in _preconditions("商城", "de-DE", "已登录并充值") or "语言" in _preconditions("商城", "de-DE", "已登录并充值")
    assert build_rule_cases(split_requirements(""), source_lang="en-US", target_lang="de-DE") == []
    a = split_requirements("【商城】\n1. 商店展示礼包标题与价格\n")
    cs = build_rule_cases(a, source_lang="en-US", target_lang="de-DE", categories=["TEXT_OVERFLOW"], max_cases=1)
    assert len(cs) <= 1 and cs[0]["source"] == "rules"


# ---------------- analysis ----------------
def test_analysis_branches():
    assert split_requirements("").points == []
    assert split_requirements("【商城系统】\n第1章 概述\n").points == []
    r = split_requirements("【商城系统】\n\n- 商店展示限时礼包标题与价格\n\n1. 商店展示限时礼包价格\n")
    assert len(r.points) >= 1
    many = "\n".join("- 需求点条目编号%d号内容足够长" % i for i in range(100))
    assert len(split_requirements(many, max_points=10).points) <= 10
    mods = split_requirements("商城展示标题\n商城展示按钮\n").modules
    assert len(mods) == len(set(mods))


# ---------------- api_server ----------------
def test_api_meta_health_library():
    m = client.get("/api/meta").json()
    assert m["limits"]["max_count"] == 2000 and m["defaults"]["count"] == 0
    assert m["limits"]["max_upload_bytes"] > 0
    assert client.get("/api/health").json()["ok"] is True
    lib = client.get("/api/case-library").json()
    assert "files" in lib and "count" in lib and "root" in lib


def test_api_parse_paths():
    ok = client.post("/api/parse", files={"file": ("需求.md", "商城展示".encode("utf-8"), "text/markdown")})
    assert ok.status_code == 200 and "商城" in ok.json()["text"]
    assert client.post("/api/parse", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 400
    orig = api_server.MAX_UPLOAD_BYTES
    try:
        api_server.MAX_UPLOAD_BYTES = 8
        r = client.post("/api/parse", files={"file": ("a.txt", b"0123456789", "text/plain")})
        assert r.status_code == 413
    finally:
        api_server.MAX_UPLOAD_BYTES = orig


def test_api_analyze_generate_paths():
    assert client.post("/api/analyze", json={"text": "   "}).status_code == 400
    a = client.post("/api/analyze", json={"text": "【商城】\n1. 商店展示礼包标题与价格\n"}).json()
    assert a["points"] and a["summary"]
    assert client.post("/api/generate", json={"text": " "}).status_code == 400
    assert client.post("/api/generate", json={"text": "【商城】\n1. 商店展示礼包标题与价格\n",
                                             "use_ai": True, "api_key": ""}).status_code == 400
    g = client.post("/api/generate", json={"text": "【商城】\n1. 商店展示礼包标题与价格\n",
                                           "use_ai": False, "count": 0}).json()
    assert g["generator"] == "rules" and g["cases"] and g["gate"]["count"] == len(g["cases"])


def test_api_export_save_settings_paths():
    cases = client.post("/api/generate", json={"text": "【商城】\n1. 商店展示礼包标题与价格\n",
                                              "count": 0}).json()["cases"]
    assert client.post("/api/export", json={"cases": [], "task_name": "空"}).status_code == 400
    r = client.post("/api/export", json={"cases": cases, "task_name": '冒烟<>:"'})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.post("/api/save", json={"cases": [], "task_name": "空"}).status_code == 400
    tmp = Path(tempfile.mkdtemp())
    orig_out = api_server.OUT_ROOT
    try:
        api_server.OUT_ROOT = tmp
        s = client.post("/api/save", json={"cases": cases, "task_name": "smoke"}).json()
        assert Path(s["folder"]).exists() and s["count"] == len(cases)
    finally:
        api_server.OUT_ROOT = orig_out
    t = client.post("/api/settings/test", json={"api_key": ""}).json()
    assert t["ok"] is False and "未配置" in t["message"]
    assert client.post("/api/settings/save", json={"api_key": "sk-x"}).json()["persisted"] is False
    assert client.post("/api/settings/save", json={"persist": True, "api_key": ""}).status_code == 400
    tmp2 = Path(tempfile.mkdtemp())
    orig_env = api_server.ENV_PATH
    try:
        api_server.ENV_PATH = tmp2 / ".env"
        saved = client.post("/api/settings/save", json={"persist": True, "provider": "openai",
                                                        "model": "gpt-4o-mini",
                                                        "base_url": "https://api.openai.com/v1",
                                                        "api_key": "sk-abcdefghijkl"}).json()
        assert saved["persisted"] is True
        assert api_server.ENV_PATH.exists()
        assert "LQA_API_KEY=sk-abcdefghijkl" in api_server.ENV_PATH.read_text(encoding="utf-8")
    finally:
        api_server.ENV_PATH = orig_env
    assert client.post("/api/shutdown").status_code == 400


def test_api_skills_endpoints():
    key = "wb-contract"
    client.delete("/api/skills/" + key)
    c = client.post("/api/skills", json={"name": "白盒技能", "body": "要点", "skill_key": key})
    assert c.status_code in (200, 201)
    assert client.get("/api/skills/" + key).status_code == 200
    assert client.put("/api/skills/" + key, json={"name": "改名"}).json()["body"] == "要点"
    assert client.get("/api/skills/nope-not-exist").status_code == 404
    assert client.put("/api/skills/nope-not-exist", json={"name": "x"}).status_code == 404
    dup = client.post("/api/skills", json={"name": "白盒技能", "body": "要点", "skill_key": key})
    assert dup.status_code == 409
    assert client.post("/api/skills", json={"name": " ", "body": "b"}).status_code == 422
    assert client.post("/api/skills", json={"name": "n", "body": ""}).status_code == 422
    assert client.delete("/api/skills/" + key).json()["removed"] is True
    assert client.delete("/api/skills/" + key).json()["removed"] is False


def test_skills_read_one_non_dict():
    tmp = Path(tempfile.mkdtemp())
    (tmp / "list.json").write_text("[1, 2, 3]", encoding="utf-8")
    (tmp / "str.json").write_text('"just a string"', encoding="utf-8")
    (tmp / "num.json").write_text("42", encoding="utf-8")
    assert sk.list_skills(tmp) == []
    for name, body in (("", "b"), ("n", ""), ("x" * 61, "b"), ("n", "y" * 4001)):
        try:
            sk._validate(name, body)
            raise AssertionError("应拒绝 name=%r body_len=%d" % (name[:5], len(body)))
        except ValueError:
            pass
    assert sk.skills_root(tmp) == tmp


def test_api_settings_save_updates_existing_env():
    tmp = Path(tempfile.mkdtemp())
    env = tmp / ".env"
    env.write_text("# keep me\nLQA_MODEL=old-model\nLQA_API_KEY=sk-oldoldold\n", encoding="utf-8")
    orig_env = api_server.ENV_PATH
    try:
        api_server.ENV_PATH = env
        first = client.post("/api/settings/save", json={"persist": True, "api_key": "sk-newkey123456",
                                                        "model": "new-model"}).json()
        assert first["persisted"] is True
        text = env.read_text(encoding="utf-8")
        assert "# keep me" in text
        assert "LQA_MODEL=new-model" in text and "LQA_API_KEY=sk-newkey123456" in text
        client.post("/api/settings/save", json={"persist": True, "api_key": "sk-thirdkey7890"})
        text2 = env.read_text(encoding="utf-8")
        assert text2.count("LQA_API_KEY") == 1 and "sk-thirdkey7890" in text2
        client.post("/api/settings/save", json={"persist": True, "api_key": "sk-fourthkey1234",
                                                "provider": "openai",
                                                "base_url": "https://api.openai.com/v1"})
        text3 = env.read_text(encoding="utf-8")
        assert "LQA_PROVIDER=openai" in text3 and "LQA_BASE_URL=https://api.openai.com/v1" in text3
    finally:
        api_server.ENV_PATH = orig_env


# ---------------- 规格探针（失败 = 已确认缺陷） ----------------
def test_probe_skills_illegal_id_is_fail_soft():
    """探针：skills 模块承诺「单个损坏文件不影响其余技能读取」。
    id 非法的技能文件也应被跳过；当前 _read_one 在 try 之外调用 _ensure_safe_id → ValueError 外泄。"""
    tmp = Path(tempfile.mkdtemp())
    (tmp / "evil.json").write_text(json.dumps({"id": "..", "name": "x", "body": "b"}), encoding="utf-8")
    (tmp / "ok.json").write_text(json.dumps({"id": "ok", "name": "n", "body": "b"}), encoding="utf-8")
    assert [s["id"] for s in sk.list_skills(tmp)] == ["ok"]


def test_probe_env_file_reload_updates_process_env():
    """探针：改写 .env 后再加载应让新值在进程内生效（密钥轮换场景）。"""
    tmp = Path(tempfile.mkdtemp())
    env = tmp / ".env"
    os.environ.pop("LQA_WB_ROTATE", None)
    env.write_text("LQA_WB_ROTATE=v1\n", encoding="utf-8")
    load_env_file(env)
    assert os.environ["LQA_WB_ROTATE"] == "v1"
    env.write_text("LQA_WB_ROTATE=v2\n", encoding="utf-8")
    load_env_file(env)
    assert os.environ["LQA_WB_ROTATE"] == "v2"


def test_probe_parse_txt_decodes_gbk():
    """探针：中文 Windows 下 GBK 编码的 .txt 应被正确解码，而非 UTF-8 replace 成乱码。"""
    assert parse_txt("商店展示限时礼包".encode("gbk")) == "商店展示限时礼包"


def test_probe_analysis_max_points_applies_to_headings():
    """探针：max_points 应对标题分支同样生效，否则纯标题文档会突破上限。"""
    text = "\n".join("# 需求标题条目编号%d内容" % i for i in range(30))
    assert len(split_requirements(text, max_points=5).points) <= 5


if __name__ == "__main__":
    g = dict(globals())
    names = sorted(n for n in g if n.startswith("test_") and callable(g[n]))
    passed = failed = 0
    for n in names:
        try:
            g[n]()
            passed += 1
            print("PASS", n)
        except Exception as e:
            failed += 1
            print("FAIL", n, "->", type(e).__name__ + ": " + str(e)[:200])
    print("WHITEBOX PASS=%d FAIL=%d TOTAL=%d" % (passed, failed, len(names)))
