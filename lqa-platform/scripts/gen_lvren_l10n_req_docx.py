# -*- coding: utf-8 -*-
"""生成「旅人」干员本地化需求 Word：技能释放实机图 + 需求表。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ASSETS = Path(r"E:\智谱\本地化测试平台\lqa-platform\assets\delta_force\lvren")
OUT = Path(r"C:\Users\zhang\Desktop\三角洲_干员旅人_本地化需求_v1.0.docx")

INK = RGBColor(0x1F, 0x1F, 0x1F)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x59, 0x59, 0x59)
FONT_CJK = "微软雅黑"


def set_run_font(run, size=11, bold=False, color=INK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), FONT_CJK)


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    rPr = normal._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(5)
    for n, size in [(1, 16), (2, 13)]:
        st = doc.styles[f"Heading {n}"]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        rPr = st._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), FONT_CJK)
    t = doc.styles["Title"]
    t.font.size = Pt(20)
    t.font.bold = True
    t.font.color.rgb = NAVY
    rPr = t._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), FONT_CJK)


def add_page_number(p):
    run = p.add_run()
    for tag, attr, text in [
        ("w:fldChar", "begin", None),
        ("w:instrText", None, "PAGE"),
        ("w:fldChar", "end", None),
    ]:
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        if text:
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_run_font(run, size=9, color=MUTED)


def header_footer(doc):
    s = doc.sections[0]
    h = s.header.paragraphs[0]
    r = h.add_run("三角洲行动 · 干员「旅人」本地化需求 · 内部资料")
    set_run_font(r, size=9, color=MUTED)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = f.add_run("第 ")
    set_run_font(r2, size=9, color=MUTED)
    add_page_number(f)
    r3 = f.add_run(" 页")
    set_run_font(r3, size=9, color=MUTED)


def para(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, bold=style.startswith("Heading"))
    return p


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, name in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(name)
        set_run_font(r, size=9, bold=True)
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            r = c.paragraphs[0].add_run(str(val) if val != "" else " ")
            set_run_font(r, size=9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def img(doc, path: Path, caption: str, width_cm=16.0):
    if not path.is_file():
        para(doc, f"[图缺失] {path.name}")
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=9, color=MUTED)


def build():
    doc = Document()
    setup(doc)
    header_footer(doc)

    t = doc.add_paragraph("干员「旅人」本地化需求", style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("三角洲行动 · S11 支援位 · 策划提测 v1.0")
    set_run_font(r, size=12, bold=True, color=MUTED)
    doc.add_paragraph()

    table(
        doc,
        ["字段", "内容"],
        [
            ["文档编号", "DF-OP-LVREN-LOC-001"],
            ["干员代号", "旅人"],
            ["本名", "罗温·沃尔什（Rowan Walsh）"],
            ["定位", "支援位 / 特殊医疗支援（可双向干扰）"],
            ["赛季", "S11"],
            ["搭档", "军犬「四叶」"],
            ["提出方", "系统策划组 / 张策"],
            ["接收方", "本地化、LQA、客户端 UI"],
            ["实机参考", "B站 BV1cGY966Eov《三角洲最特殊的支援位——一期视频教会你旅人》"],
            ["日期", "2026-09-18"],
            ["优先级", "P0 — 阻塞 S11 海外包"],
        ],
        [3.2, 12.3],
    )
    para(
        doc,
        "说明：本文档只覆盖干员「旅人」及其军犬「四叶」的可本地化文案（技能名、技能描述、状态提示、局内语音、HUD）。不覆盖武器通用文案与全量 UI。技能表现以实机释放画面为准。",
    )

    # ---- 干员档案 ----
    doc.add_paragraph("1. 干员档案与形象", style="Heading 1")
    para(doc, "旅人为攀登/救援背景的支援干员，标志性外观：护目镜、攀登绳、肩扛气雾针剂枪，常与军犬四叶同框。译名需保留「远行/旅人」意象，禁止译成单纯旅行者。")

    img(doc, ASSETS / "cover_deep.jpg", "图1  干员立绘参考：旅人 + 军犬四叶（来源：B站精讲封面，仅作形象对照）")
    img(doc, ASSETS / "cover_syringe.jpg", "图2  官方宣发图：S11 旅人气雾针剂枪「双向机制」")

    table(
        doc,
        ["中文", "英文锁定（建议）", "本地化要求"],
        [
            ["旅人", "Wayfarer / Voyager（待法务与发行拍板）", "全语言统一；禁止 Wanderer（语气过轻）"],
            ["罗温·沃尔什", "Rowan Walsh", "音译一致；档案/干员详情页用全名"],
            ["四叶", "Clover", "军犬名，不意译成 Four Leaves；首字母大写"],
            ["支援位", "Support", "职位标签，与突击/侦察/工程并列"],
            ["气雾针剂枪", "Aerosol Syringe Gun", "装备名，全技能描述一致"],
        ],
        [2.8, 5.5, 7.2],
    )

    # ---- 技能总览 ----
    doc.add_paragraph("2. 技能组总览（本地化范围）", style="Heading 1")
    para(doc, "下列技能名取自实机 HUD 顶栏（见图3）。所有技能名、短描述、长描述、冷却提示、局内喊话均须本地化。")

    img(
        doc,
        ASSETS / "skill_deep_005.jpg",
        "图3  实机技能 HUD 顶栏：紧急勤务 | 气雾针剂枪 | 刺激性 | 军犬协同 | 补充",
    )

    table(
        doc,
        ["槽位", "技能名", "机制摘要（策划）", "文案类型", "优先级"],
        [
            ["被动/特质", "紧急勤务", "倒地/紧急状态相关支援增益；技能说明面板需完整可读", "技能名+长短描述+图标 tooltip", "P0"],
            ["主武器技", "气雾针剂枪", "可对队友回血、对敌造成气雾干扰（双向）；局内切换提示", "技能名+模式切换提示+命中反馈", "P0"],
            ["战术/状态", "刺激性", "刺激性气雾：敌方接触后出现咳嗽等 debuff，可被侦测", "技能名+状态名+状态描述", "P0"],
            ["搭档技", "军犬协同", "放出四叶协同作战；短按可指派位置；击杀播报含犬名", "技能名+指令提示+播报模板", "P0"],
            ["补给", "补充", "弹药/针剂补给相关；进度条与完成提示", "技能名+进度文案", "P1"],
        ],
        [2.2, 2.4, 5.0, 3.5, 1.4],
    )

    # ---- 气雾针剂枪 ----
    doc.add_paragraph("3. 技能：气雾针剂枪（核心）", style="Heading 1")
    para(doc, "该技能为旅人识别度最高的释放技能。本地化必须覆盖：持枪第一人称、对敌喷射气雾、对队友注射/喷雾的全部 UI 与状态文案。")

    img(
        doc,
        ASSETS / "skill_main_032.jpg",
        "图4  释放气雾后：敌方进入气雾区域并触发「咳嗽」状态提示（实机）",
    )
    img(
        doc,
        ASSETS / "skill_guide_010.jpg",
        "图5  教程说明：敌人进入烟会咳嗽，可帮助判断位置（半辅助半侦察定位）",
    )

    para(doc, "3.1 必译 UI 字符串（节选）", style="Heading 2")
    table(
        doc,
        ["上下文", "源文案（中文示例）", "要求", "LQA 检查点"],
        [
            ["技能名", "气雾针剂枪", "全界面一致，禁拆词", "术语一致"],
            ["模式提示", "当前为治疗模式 / 当前为干扰模式", "短句，≤12 汉字等效", "溢出"],
            ["对敌状态", "刺激性气雾", "与状态栏名称一致", "术语"],
            ["状态效果", "咳嗽 / 视野受限 / 可被侦测", "可判定，禁「异常状态」", "可读性"],
            ["命中反馈", "已治疗 {ally} / 已标记 {enemy}", "占位符不翻译", "占位符"],
            ["冷却/充能", "针剂充能中 {time}", "时间格式 locale 化", "日期数字"],
        ],
        [2.8, 4.2, 4.5, 4.0],
    )

    # ---- 刺激性 ----
    doc.add_paragraph("4. 技能：刺激性（状态与侦测）", style="Heading 1")
    img(
        doc,
        ASSETS / "skill_main_032.jpg",
        "图6  刺激性气雾生效：敌方接触后「开始咳嗽」——此状态名必须进术语表（同图4场景，强调状态文案）",
    )
    table(
        doc,
        ["条目", "说明", "本地化红线"],
        [
            ["状态名：咳嗽", "敌方进入气雾后的核心 debuff", "禁止译成笼统「中毒」，需保留刺激/呛咳语义"],
            ["状态描述", "接触刺激性气雾，发出咳嗽并短暂暴露位置", "描述与实测一致；德/日注意长句换行"],
            ["己方提示", "侦测到咳嗽信号", "HUD 短句，击杀/辅助播报可复用"],
            ["音效字幕（如有）", "咳嗽声", "与语音字幕轨对齐"],
        ],
        [3.0, 5.5, 7.0],
    )

    # ---- 军犬 ----
    doc.add_paragraph("5. 技能：军犬协同（搭档四叶）", style="Heading 1")
    para(doc, "四叶为独立可命名单位，局内 HUD、击杀播报、指令提示均含犬名。本地化不得把「四叶」当普通名词翻译。")

    img(
        doc,
        ASSETS / "skill_deep_035.jpg",
        "图7  军犬协同作战中的局内 HUD（四叶持续作战提示）",
    )

    table(
        doc,
        ["场景", "源文案示例", "要求"],
        [
            ["技能名", "军犬协同", "与技能栏一致"],
            ["指令短按", "短按 {key} 指派四叶前往指定位置", "键位名不翻译；犬名统一 Clover"],
            ["击杀播报", "四叶 击倒了 {enemy}", "播报极短，防截断；阿语 RTL 单独验"],
            ["犬只状态", "四叶：作战中 / 召回中 / 受伤", "状态枚举，禁同义多译"],
            ["档案", "军犬四叶与旅人长期搭档…", "可叙事；专有名词与术语表一致"],
        ],
        [2.8, 5.5, 7.2],
    )

    # ---- 语音 ----
    doc.add_paragraph("6. 局内语音与字幕", style="Heading 1")
    table(
        doc,
        ["触发", "中文示例", "本地化要求"],
        [
            ["释放气雾", "「气雾已部署，注意呼吸。」", "语气：冷静支援，非口号"],
            ["治疗队友", "「稳住，给你一针。」", "口语但专业；禁网络用语"],
            ["放出四叶", "「四叶，去！」", "犬名锁定；短促"],
            ["侦测咳嗽", "「他们中招了。」", "可与 HUD 提示呼应"],
            ["倒地求救", "「我需要支援……」", "P0 字幕与音频对齐 LQA"],
        ],
        [2.8, 5.0, 7.7],
    )

    # ---- LQA ----
    doc.add_paragraph("7. LQA 验收标准（旅人专项）", style="Heading 1")
    table(
        doc,
        ["检查项", "合格标准", "严重级别", "对应实机画面"],
        [
            ["技能名一致性", "紧急勤务/气雾针剂枪/刺激性/军犬协同/补充 五处译名全局一致", "Critical", "图3"],
            ["气雾状态名", "「咳嗽」及描述完整显示，无截断", "Critical", "图4/6"],
            ["双向模式提示", "治疗/干扰切换文案即时且不重叠", "Critical", "图4"],
            ["四叶命名", "全语言为 Clover 锁定译法，播报不截断", "Critical", "图7"],
            ["占位符", "{ally}/{enemy}/{key}/{time} 运行时替换", "Critical", "技能表"],
            ["字幕对齐", "释放语音与字幕同句同时长", "Major", "第6章"],
            ["德语扩展", "技能描述面板不截断（气雾针剂枪复合词）", "Major", "图3 tooltip"],
            ["阿语 RTL", "技能栏与状态图标镜像正确", "Critical", "图3"],
        ],
        [2.6, 6.5, 2.0, 4.4],
    )

    # ---- 参考 ----
    doc.add_paragraph("8. 素材与参考来源", style="Heading 1")
    table(
        doc,
        ["类型", "来源", "用途"],
        [
            ["实机演示（主参考）", "B站 BV1cGY966Eov", "旅人支援位实战与技能释放"],
            ["技能精讲", "B站 BV1VrbZ6aEZ5", "HUD 技能栏与局内作战帧"],
            ["针剂枪机制", "B站 BV1GebK6LEVy", "气雾针剂枪双向机制"],
            ["教程", "B站 BV17sbK6HEVF", "咳嗽 debuff 与定位说明"],
            ["宣发图", "S11 气雾针剂枪双向机制图", "形象与装备名"],
        ],
        [3.5, 5.0, 7.0],
    )
    para(doc, "注：B 站分帧预览分辨率有限（480×270），正式 LQA 请以海外测试包实机截屏替换本文档插图；技能名与文案条目已按实机 HUD 固化。")

    doc.add_paragraph("9. 交付与责任", style="Heading 1")
    table(
        doc,
        ["交付物", "负责", "计划日期"],
        [
            ["旅人技能术语表（中/英/目标语）", "本地化", "2026-09-25"],
            ["技能 UI + 状态文案包", "本地化", "2026-10-05"],
            ["局内语音字幕", "本地化", "2026-10-08"],
            ["旅人专项 LQA 用例（≥15 条）", "QA", "2026-10-12"],
            ["首轮缺陷报告", "QA", "2026-10-18"],
        ],
        [6.0, 3.5, 6.0],
    )

    doc.add_paragraph()
    para(doc, "—— 以下无正文 ——")
    doc.core_properties.title = "三角洲行动干员旅人本地化需求"
    doc.core_properties.author = "系统策划组 / 张策"
    doc.core_properties.subject = "S11 支援位干员旅人技能本地化"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    if OUT.exists():
        alt = OUT.with_name(OUT.stem + "_alt.docx")
        doc.save(alt)
        print(f"OK (alt) {alt}")
    else:
        doc.save(OUT)
        print(f"OK {OUT}")


if __name__ == "__main__":
    build()
