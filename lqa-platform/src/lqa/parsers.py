# -*- coding: utf-8 -*-
"""需求文档解析：docx / xlsx / md / txt / 粘贴文本 → 统一纯文本。"""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_text(raw: bytes) -> str:
    """字节 → 文本：优先 UTF-8（含 BOM），回落 GBK（中文 Windows 存量 .txt）。

    两者均严格解码，都失败才退化为 UTF-8 replace，避免把 GBK 文件
    解成乱码后静默进入后续拆解流程。
    """
    for enc in ("utf-8-sig", "gbk"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_txt(raw: str | bytes) -> str:
    if isinstance(raw, bytes):
        raw = _decode_text(raw)
    return _normalize(raw)


def parse_md(raw: str | bytes) -> str:
    return parse_txt(raw)


def parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for ti, table in enumerate(doc.tables, 1):
        parts.append(f"\n[表格{ti}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _normalize("\n".join(parts))


def parse_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    parts: list[str] = []
    for name in wb.sheetnames:
        ws = wb[name]
        parts.append(f"[Sheet] {name}")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > 500:
                break
            vals = ["" if v is None else str(v).strip() for v in row]
            if any(vals):
                parts.append(" | ".join(vals))
    wb.close()
    return _normalize("\n".join(parts))


def parse_upload(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in (".md", ".markdown"):
        return parse_md(data)
    if ext in (".txt", ".text", ".log"):
        return parse_txt(data)
    if ext in (".docx",):
        return parse_docx(data)
    if ext in (".xlsx", ".xlsm"):
        return parse_xlsx(data)
    if ext in (".doc", ".xls"):
        raise ValueError(
            f"暂不支持旧版 {ext}，请另存为 .docx / .xlsx 后再上传"
        )
    raise ValueError(f"不支持的文件类型: {ext or '(无扩展名)'}")


def merge_sources(doc_text: str, paste_text: str) -> str:
    a = _normalize(doc_text or "")
    b = _normalize(paste_text or "")
    if a and b:
        return a + "\n\n" + b
    return a or b
